"""
Build an ATS-safe cover letter DOCX from a JSON description.

Usage
-----
    python build_cover_letter_docx.py --input cover.json --output CoverLetter.docx

Or pipe JSON via stdin:
    cat cover.json | python build_cover_letter_docx.py --output CoverLetter.docx

JSON schema
-----------
{
  "name": "Jane Doe",
  "contact": {
    "address": "Street Address, City, State Zip",
    "phone": "555-555-5555",
    "email": "jane@example.com",
    "linkedin": "linkedin.com/in/janedoe"
  },
  "date": "May 6, 2026",
  "recipient": {
    "name": "Hiring Manager",
    "title": "",
    "company": "Acme Corp",
    "address": ""
  },
  "salutation": "Dear Hiring Manager,",
  "paragraphs": [
    "Opening paragraph...",
    "Experience paragraph...",
    "Differentiator paragraph...",
    "Closing paragraph..."
  ],
  "signoff": "Sincerely,"
}
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).resolve().parent))
from verify_output import (  # noqa: E402
    build_report,
    find_config,
    load_config,
    verify_anonymization,
    verify_page_count,
)


COVER_LETTER_MAX_PAGES = 1


BODY_FONT = "Calibri"
BODY_SIZE = Pt(11)
NAME_SIZE = Pt(16)

# ATS safety: Unicode dashes (en-dash U+2013, em-dash U+2014) render as
# garbage in some ATS parsers. Defensive substitution before any text hits
# the document.
_DASH_NORMALIZE = str.maketrans({"–": "-", "—": "-"})


def _ats_safe(text):
    if text is None:
        return text
    if isinstance(text, str):
        return text.translate(_DASH_NORMALIZE)
    return text


def _normalize_strings(obj):
    if isinstance(obj, str):
        return _ats_safe(obj)
    if isinstance(obj, list):
        return [_normalize_strings(x) for x in obj]
    if isinstance(obj, dict):
        return {k: _normalize_strings(v) for k, v in obj.items()}
    return obj


def _set_run_font(run, size=BODY_SIZE, bold=False, italic=False):
    run.font.name = BODY_FONT
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)


def _set_paragraph_spacing(paragraph, before=0, after=8, line_spacing=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing


def _add_line(doc, text, *, bold=False, italic=False, size=BODY_SIZE, before=0, after=0, line_spacing=1.15):
    p = doc.add_paragraph()
    _set_paragraph_spacing(p, before=before, after=after, line_spacing=line_spacing)
    if text:
        run = p.add_run(text)
        _set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def build_cover_letter(data: dict, output_path: Path) -> Path:
    data = _normalize_strings(data)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE

    # Letterhead — name + contact
    _add_line(doc, data["name"], bold=True, size=NAME_SIZE, after=0)
    contact = data.get("contact", {})
    contact_parts = [
        contact.get("address", ""),
        contact.get("phone", ""),
        contact.get("email", ""),
        contact.get("linkedin", ""),
    ]
    contact_line = " · ".join([p for p in contact_parts if p])
    _add_line(doc, contact_line, size=Pt(10), after=12)

    # Date
    if data.get("date"):
        _add_line(doc, data["date"], after=12)

    # Recipient
    recipient = data.get("recipient", {})
    if recipient:
        if recipient.get("name"):
            _add_line(doc, recipient["name"], after=0)
        if recipient.get("title"):
            _add_line(doc, recipient["title"], after=0)
        if recipient.get("company"):
            _add_line(doc, recipient["company"], after=0)
        if recipient.get("address"):
            _add_line(doc, recipient["address"], after=0)
        # spacer
        _add_line(doc, "", after=8)

    # Salutation
    _add_line(doc, data.get("salutation", "Dear Hiring Manager,"), after=10)

    # Body paragraphs
    for para in data.get("paragraphs", []):
        _add_line(doc, para, after=10, line_spacing=1.15)

    # Signoff
    _add_line(doc, data.get("signoff", "Sincerely,"), before=4, after=24)
    _add_line(doc, data["name"], after=0)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n", 1)[0])
    parser.add_argument("--input", "-i", help="Path to JSON. If omitted, reads stdin.")
    parser.add_argument("--output", "-o", required=True, help="Output DOCX path.")
    parser.add_argument(
        "--config",
        type=Path,
        default=None,
        help=(
            "Path to tailor_config.json. If omitted, walks up from the output "
            "directory looking for the file."
        ),
    )
    parser.add_argument(
        "--skip-anonymization",
        action="store_true",
        help=(
            "Skip the anonymization check (use when the target JD employer matches "
            "one of your configured anonymization patterns)."
        ),
    )
    parser.add_argument(
        "--no-strict-pages",
        action="store_true",
        help="Warn but don't fail on page-count overage. Default cap is 1 page.",
    )
    parser.add_argument(
        "--skip-page-check",
        action="store_true",
        help="Skip the page-count check entirely (e.g. when MS Word isn't available).",
    )
    args = parser.parse_args(argv)

    if args.input:
        data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    else:
        data = json.loads(sys.stdin.read())

    out = build_cover_letter(data, Path(args.output))

    config_path = args.config or find_config([out.parent, Path.cwd()])
    config = load_config(config_path)

    violations = verify_anonymization(out, config=config, skip=args.skip_anonymization)
    page_count: int | None = None
    if not args.skip_page_check:
        page_count = verify_page_count(out, max_pages=COVER_LETTER_MAX_PAGES)

    report = build_report(
        out,
        violations,
        page_count,
        COVER_LETTER_MAX_PAGES if not args.skip_page_check else None,
        config_path,
    )
    exit_code = 0

    if violations:
        exit_code |= 1
        print(
            f"\nERROR: rendered cover letter {out.name} contains "
            f"{len(violations)} anonymization violation(s):",
            file=sys.stderr,
        )
        for v in violations:
            print(f"  - {v.token!r} matched {v.match!r} near: {v.context!r}", file=sys.stderr)
        print(
            "\nFix the cover letter JSON to use anonymized framing per your master_profile.md, "
            "then re-run. Pass --skip-anonymization only when the target JD employer matches "
            "one of your configured anonymization patterns.",
            file=sys.stderr,
        )

    if report["page_over_cap"]:
        level = "WARN" if args.no_strict_pages else "ERROR"
        print(
            f"\n{level}: rendered cover letter {out.name} is {page_count} pages "
            f"(cap is {COVER_LETTER_MAX_PAGES}).",
            file=sys.stderr,
        )
        if not args.no_strict_pages:
            exit_code |= 2
            print(
                "Tighten to 2-3 paragraphs / ~250-300 words. See references/cover_letter_guide.md.",
                file=sys.stderr,
            )

    print(json.dumps(report))
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
